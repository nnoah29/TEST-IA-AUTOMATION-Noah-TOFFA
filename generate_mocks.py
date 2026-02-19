
import os
import csv
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw

OUTPUT_DIR = "data/fanga_inbox"

def create_mocks():
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    files = [
        ("contrat_aissata_kone_2024.pdf", "pdf"),
        ("facture_station_cocody_mars.pdf", "pdf"),
        ("photo_station_plateau_01.jpg", "image"),
        ("rapport_mensuel_conducteurs.xlsx", "excel"),
        ("export_transactions_fevrier.csv", "csv"),
        ("carte_identite_yacouba.png", "image"),
        ("maintenance_batterie_ST-002.docx", "word"),
        ("planning_equipe_avril.pdf", "pdf"),
        ("bon_de_commande_motos.pdf", "pdf"),
        ("screenshot_app_bug.png", "image"),
    ]

    for filename, type in files:
        filepath = os.path.join(OUTPUT_DIR, filename)
        
        if type == "pdf":
            c = canvas.Canvas(filepath)
            c.drawString(100, 750, f"Contenu de {filename}")
            c.save()
            print(f"Created PDF: {filename}")
            
        elif type == "image":
            img = Image.new('RGB', (100, 100), color = 'red')
            d = ImageDraw.Draw(img)
            d.text((10,10), filename, fill=(255,255,0))
            img.save(filepath)
            print(f"Created Image: {filename}")
            
        elif type == "word":
            doc = Document()
            doc.add_paragraph(f"Contenu de {filename}")
            doc.save(filepath)
            print(f"Created Word: {filename}")
            
        elif type == "excel":
            wb = Workbook()
            ws = wb.active
            ws['A1'] = f"Contenu de {filename}"
            wb.save(filepath)
            print(f"Created Excel: {filename}")
            
        elif type == "csv":
            with open(filepath, 'w', newline='') as csvfile:
                writer = csv.writer(csvfile)
                writer.writerow(['Colonne1', 'Colonne2'])
                writer.writerow([f"Contenu de {filename}", 'Data'])
            print(f"Created CSV: {filename}")

if __name__ == "__main__":
    create_mocks()
